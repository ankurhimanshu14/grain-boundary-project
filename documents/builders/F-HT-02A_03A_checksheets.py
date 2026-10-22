from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from datetime import datetime

F="Arial"
def font(sz=10,b=False,color="000000",i=False): return Font(name=F,size=sz,bold=b,color=color,italic=i)
thin=Side(style="thin",color="8A9AA0"); BOX=Border(left=thin,right=thin,top=thin,bottom=thin)
INPUT=PatternFill("solid",fgColor="FFF6C8")   # cells to fill in
HEAD=PatternFill("solid",fgColor="1F4E5A"); SEC=PatternFill("solid",fgColor="DCE6E8"); CALC=PatternFill("solid",fgColor="F2F4F4")
WRAP=Alignment(wrap_text=True,vertical="top"); CEN=Alignment(horizontal="center",vertical="center",wrap_text=True)
BLUE="0000FF"
FURN=["F-01","F-02","F-03","F-04"]
DT='dd-mmm-yyyy hh:mm'

START=[
 ("A. Before ignition",None),
 ("Maintenance work closed, permits returned, guards and covers fitted","All",1,""),
 ("Gas supply pressure checked (record the value, mbar)","All",1,"mbar"),
 ("Combustion air blower running; burner safety interlocks and flame detectors tested","All",1,""),
 ("Furnace purged before ignition by the burner management sequence","All",1,""),
 ("Control thermocouples TC-01 … TC-06 connected in both chambers; both walls of every zone read plausibly (odd = right wall, even = left wall)","All",1,""),
 ("SAT of the control thermocouples in date (CALOR → Pyrometry → SAT)","All",1,""),
 ("Over-temperature protection set and active","All",1,""),
 ("SCADA recording running and its clock correct","All",1,""),
 ("B. Equipment",None),
 ("Pusher, extractor, doors and rail transfer tested empty","All",0,""),
 ("Controlled cooling chamber: air blower, rail transfer in / out and cooling timer working","All",0,""),
 ("Quench tank (water): level within the marks, water temperature within the limit, agitation running (checked by eye: agitation is not monitored)","F-01",1,"°C"),
 ("Quench tank (oil): level within the marks, oil temperature 50–80 °C, agitation and oil cooler running (checked by eye), no water / crackle test as scheduled","F-02",1,"°C"),
 ("Exhaust / flue draught and door curtains working","All",0,""),
 ("Trays for the first batches inspected: HK trays without cracks or distortion; load per tray not above 200 kg","All",0,""),
 ("C. Heat-up and release",None),
 ("Heat-up rate of the furnace manual followed (refractory)","All",0,""),
 ("Set values entered from the process specification in force for the first part (record the specification no.)","All",1,"Spec no."),
 ("All zones within tolerance and soaked before the first tray (times and zone readings below)","All",1,""),
 ("First-off hardness test of the first batch done within the time below (record the test no.)","All",1,"Test no."),
 ("Start-up recorded in CALOR → Process → Alarms (start-up record)","All",0,""),
]
SHUT=[
 ("A. Before stopping",None),
 ("Production stop agreed and the reason recorded above","All",1,""),
 ("Last batch identified in the daily production report (part, batch, last tray no.)","All",1,"Last tray no."),
 ("Loading stopped; empty trays pushed to run the line out","All",1,""),
 ("Line emptied of product, or the product left inside recorded below","All",1,""),
 ("B. Stopping",None),
 ("Burners shut down by the burner management sequence","All",1,""),
 ("Gas isolation valves closed (and locked where required)","All",1,""),
 ("Combustion air / exhaust left running as the furnace manual requires","All",0,""),
 ("Controlled cooling chamber empty; air blower off","All",0,""),
 ("Quench tank (water): tank covered, circulation as the procedure requires","F-01",0,""),
 ("Quench tank (oil): circulation and cooling as the procedure requires, tank covered (fire risk)","F-02",1,""),
 ("Doors closed; pusher, extractor and rail transfer in a safe position","All",0,""),
 ("SCADA recording kept running until the burners were off","All",1,""),
 ("C. After stopping",None),
 ("Hearths, rails and tray positions checked for fallen parts once cool enough (trap points)","All",1,""),
 ("Shutdown recorded in CALOR → Process → Alarms; handover note to maintenance / the next shift","All",0,""),
]

def setw(ws,widths):
    for c,w in widths.items(): ws.column_dimensions[c].width=w

def cell(ws,ref,v=None,f=None,fill=None,al=None,bd=True,fmt=None):
    c=ws[ref]
    if v is not None: c.value=v
    c.font=f or font(); 
    if fill: c.fill=fill
    c.alignment=al or Alignment(vertical="center",wrap_text=True)
    if bd: c.border=BOX
    if fmt: c.number_format=fmt
    return c

def mbox(ws,rng,v,f=None,fill=None,al=None):
    ws.merge_cells(rng)
    first=rng.split(":")[0]
    for row in ws[rng]:
        for c in row: c.border=BOX; 
        for c in row:
            if fill: c.fill=fill
    return cell(ws,first,v,f,fill,al)

def header(ws,title,docno,kind):
    mbox(ws,"A1:H1","Mass Wire and Steels Pvt. Ltd. · Bhiwadi · Heat treatment",font(10,True,"FFFFFF"),HEAD)
    mbox(ws,"A2:F3",title,font(15,True),None,Alignment(vertical="center"))
    mbox(ws,"G2:H2",f"Form {docno} · Rev 0",font(9))
    mbox(ws,"G3:H3",f"Ref.: {'WI-HT-02 Start-up' if kind=='su' else 'WI-HT-03 Shut-down'}",font(9))
    ws.row_dimensions[2].height=18; ws.row_dimensions[3].height=18

def build(ws,kind,ex=None):
    su=kind=="su"; ex=ex or {}
    title="Furnace start-up checksheet" if su else "Furnace shutdown checksheet"
    header(ws,title,"F-HT-02A" if su else "F-HT-03A",kind)
    setw(ws,{"A":5,"B":62,"C":8,"D":9,"E":11,"F":16,"G":14,"H":12,"I":12})
    # record details
    r=5
    lab=lambda ref,t: cell(ws,ref,t,font(9,True),SEC)
    inp=lambda ref,v="",fmt=None: cell(ws,ref,v if v!="" else None,font(10,False,BLUE if v!="" else "000000"),INPUT,fmt=fmt)
    # layout: label B, value C..D etc. use pairs
    rows=[("Furnace","furn",None),("Date","date","dd-mmm-yyyy"),("Shift","shift",None),
          (("Type of start-up" if su else "Type of shutdown"),"type",None),("Reason","reason",None)]
    # two columns of fields: B label / C-E value ; F label / G-H value
    pos=[(5,"B","C:E"),(5,"F","G:H"),(6,"B","C:E"),(6,"F","G:H"),(7,"B","C:H")]
    for (lbl,key,fmt),(rr,lc,vc) in zip(rows,pos):
        cell(ws,f"{lc}{rr}",lbl,font(9,True),SEC)
        a,b=vc.split(":"); mbox(ws,f"{a}{rr}:{b}{rr}",ex.get(key),font(10,False,BLUE if ex.get(key) else "000000"),INPUT)
        if fmt: ws[f"{a}{rr}"].number_format=fmt
    for rr in (5,6,7): cell(ws,f"A{rr}",None,fill=SEC)
    dvf=DataValidation(type="list",formula1='"F-01,F-02,F-03,F-04"',allow_blank=True); ws.add_data_validation(dvf); dvf.add("C5")
    dvs=DataValidation(type="list",formula1='"A (06:00-14:00),B (14:00-22:00),C (22:00-06:00)"',allow_blank=True); ws.add_data_validation(dvs); dvs.add("C6")
    types='"Cold start (from ambient),Restart after power failure,Restart after maintenance,Restart after an alarm stop,Restart after holiday / idle"' if su else '"Planned (no orders),Planned maintenance,Holiday / weekend,Emergency (alarm / power / gas),Breakdown"'
    dvt=DataValidation(type="list",formula1=types,allow_blank=True); ws.add_data_validation(dvt); dvt.add("G5")
    ws["G5"].alignment=Alignment(vertical="center",wrap_text=True)
    # settings (assumptions)
    cell(ws,"A9","Settings",font(9,True),SEC); ws.merge_cells("A9:B9")
    sets=[("Temperature tolerance ± °C",10,"tol","CALOR default ±10 °C (specification htol / ttol). Change to the tolerance of the specification in force."),
          ("Minimum soak before the first tray (min)",30,"soak","Site rule from CALOR → Alarms → Settings (soakMin 30 min); confirm against WI-HT-02 / the furnace manual."),
          ("First-off hardness test within (h)",8,"first","Site rule from CALOR → Alarms → Settings (firstHrs 8 h).")] if su else \
         [("Temperature tolerance ± °C",10,"tol","CALOR default ±10 °C.")]
    REF={}
    r=10
    for lbl,v,k,note in sets:
        mbox(ws,f"A{r}:B{r}",lbl,font(9)); c=cell(ws,f"C{r}",v,font(10,False,BLUE),INPUT,CEN); c.comment=Comment(note,"CALOR"); REF[k]=f"$C${r}"; r+=1
    r+=1
    # checklist
    hdr=["No.","Check item","Critical","Applies to","Result","Value / reference","Checked by","Time","Applicable"]
    for i,h in enumerate(hdr): cell(ws,f"{'ABCDEFGHI'[i]}{r}",h,font(9,True,"FFFFFF"),HEAD,CEN)
    ws.row_dimensions[r].height=28
    top=r+1; r=top; n=0
    items=START if su else SHUT
    exres=ex.get("res",{})
    for it in items:
        if it[1] is None:
            mbox(ws,f"A{r}:I{r}",it[0],font(10,True),SEC); r+=1; continue
        n+=1; t,ap,crit,unit=it
        cell(ws,f"A{r}",n,al=CEN); cell(ws,f"B{r}",t+(f"  [{unit}]" if unit else ""),al=WRAP)
        cell(ws,f"C{r}","Yes" if crit else "No",font(10,bool(crit),"9C0006" if crit else "000000"),al=CEN)
        cell(ws,f"D{r}",ap,al=CEN)
        e=exres.get(n,("","","",""))
        for col,val,fmt in (("E",e[0],None),("F",e[1],None),("G",e[2],None),("H",e[3],"hh:mm")):
            c=cell(ws,f"{col}{r}",val if val!="" else None,font(10,False,BLUE if val!="" else "000000"),INPUT,CEN if col in "EH" else WRAP,fmt=fmt)
        cell(ws,f"I{r}",f'=IF(D{r}="All","Yes",IF($C$5="","Yes",IF(D{r}=$C$5,"Yes","No – N/A for "&$C$5)))',font(9),CALC,CEN)
        ws.row_dimensions[r].height=30 if len(t)<70 else 42
        r+=1
    bot=r-1
    dvr=DataValidation(type="list",formula1='"OK,Not OK,N/A"',allow_blank=True); ws.add_data_validation(dvr); dvr.add(f"E{top}:E{bot}")
    # highlight Not OK, and rows not applicable
    red=PatternFill("solid",fgColor="F8D2D2"); grey=PatternFill("solid",fgColor="E6E6E6")
    ws.conditional_formatting.add(f"E{top}:E{bot}",FormulaRule(formula=[f'E{top}="Not OK"'],fill=red,font=Font(name=F,bold=True,color="9C0006")))
    ws.conditional_formatting.add(f"B{top}:H{bot}",FormulaRule(formula=[f'LEFT($I{top},2)="No"'],fill=grey,font=Font(name=F,color="808080")))
    R=f"{top}:{bot}"
    rng=lambda c:f"${c}${top}:${c}${bot}"
    r+=1
    # timing / product block
    if su:
        mbox(ws,f"A{r}:I{r}","Heat-up, soak and first tray",font(10,True,"FFFFFF"),HEAD); r+=1
        tm=[("Burners lit (heat-up started)","lit"),("All zones within tolerance (both chambers)","ready"),("First tray pushed in","push"),("First-off hardness test done","test")]
        T={}
        for lbl,k in tm:
            mbox(ws,f"A{r}:B{r}",lbl,font(9,True),SEC)
            v=ex.get(k); mbox(ws,f"C{r}:D{r}",v,font(10,False,BLUE if v else "000000"),INPUT,CEN); ws[f"C{r}"].number_format=DT
            T[k]=f"$C${r}"; r+=1
        mbox(ws,f"A{r}:B{r}","Soak before the first tray (min)",font(9,True),SEC)
        mbox(ws,f"C{r}:D{r}",f'=IF(OR({T["ready"]}="",{T["push"]}=""),"",ROUND(({T["push"]}-{T["ready"]})*1440,0))',font(10,True),CALC,CEN); SOAK=f"$C${r}"
        mbox(ws,f"E{r}:I{r}",f'=IF({SOAK}="","Enter the times above",IF({SOAK}<0,"Check the times: first tray before the zones were in tolerance",IF({SOAK}>={REF["soak"]},"OK: soak "&{SOAK}&" min (minimum "&{REF["soak"]}&")","SHORT: "&{SOAK}&" min (minimum "&{REF["soak"]}&") – product of the first trays at risk")))',font(10,True),CALC); SOAKS=f"$E${r}"; r+=1
        mbox(ws,f"A{r}:B{r}","First-off test after the first tray (h)",font(9,True),SEC)
        mbox(ws,f"C{r}:D{r}",f'=IF(OR({T["push"]}="",{T["test"]}=""),"",ROUND(({T["test"]}-{T["push"]})*24,1))',font(10,True),CALC,CEN); FT=f"$C${r}"
        mbox(ws,f"E{r}:I{r}",f'=IF({FT}="","Enter the times above",IF({FT}<0,"Check the times",IF({FT}<={REF["first"]},"OK: within "&{REF["first"]}&" h","LATE: more than "&{REF["first"]}&" h after the first tray")))',font(10,True),CALC); FTS=f"$E${r}"; r+=2
        # zone readings
        mbox(ws,f"A{r}:I{r}","Zone readings when declared ready (°C) – right wall = odd TC, left wall = even TC",font(10,True,"FFFFFF"),HEAD); r+=1
        for i,h in enumerate(["","Chamber / zone","Set value","TC right","TC left","Deviation right","Deviation left","Result",""]):
            if h: cell(ws,f"{'ABCDEFGHI'[i]}{r}",h,font(9,True),SEC,CEN)
        r+=1; z0=r
        zx=ex.get("zones",{})
        for ch in ("Hardening","Tempering"):
            for z in (1,2,3):
                tr,tl=f"TC-0{2*z-1}",f"TC-0{2*z}"
                cell(ws,f"B{r}",f"{ch} zone {z}  ({tr} right / {tl} left)")
                v=zx.get((ch,z),("","",""))
                for col,val in zip("CDE",v): cell(ws,f"{col}{r}",val if val!="" else None,font(10,False,BLUE if val!="" else "000000"),INPUT,CEN,fmt="0")
                cell(ws,f"F{r}",f'=IF(OR(C{r}="",D{r}=""),"",D{r}-C{r})',font(10),CALC,CEN,fmt='+0;-0;0')
                cell(ws,f"G{r}",f'=IF(OR(C{r}="",E{r}=""),"",E{r}-C{r})',font(10),CALC,CEN,fmt='+0;-0;0')
                cell(ws,f"H{r}",f'=IF(OR(F{r}="",G{r}=""),"",IF(AND(ABS(F{r})<={REF["tol"]},ABS(G{r})<={REF["tol"]}),"In tol.","OUT"))',font(10,True),CALC,CEN)
                r+=1
        z1=r-1
        ws.conditional_formatting.add(f"H{z0}:H{z1}",FormulaRule(formula=[f'H{z0}="OUT"'],fill=red,font=Font(name=F,bold=True,color="9C0006")))
        r+=1
    else:
        mbox(ws,f"A{r}:I{r}","Times and product left in the furnace",font(10,True,"FFFFFF"),HEAD); r+=1
        T={}
        for lbl,k in [("Loading stopped (last tray with product pushed in)","stop"),("Last tray with product out of the tempering chamber","lastout"),("Burners off","off"),("Gas isolated","gas")]:
            mbox(ws,f"A{r}:B{r}",lbl,font(9,True),SEC); v=ex.get(k)
            mbox(ws,f"C{r}:D{r}",v,font(10,False,BLUE if v else "000000"),INPUT,CEN); ws[f"C{r}"].number_format=DT; T[k]=f"$C${r}"; r+=1
        r+=1
        for i,h in enumerate(["","Product left inside","Trays","Part / batch no.","","","Disposition / reaction","",""]):
            if h: cell(ws,f"{'ABCDEFGHI'[i]}{r}",h,font(9,True),SEC,CEN)
        ws.merge_cells(f"D{r}:F{r}"); ws.merge_cells(f"G{r}:I{r}"); r+=1
        p0=r; px=ex.get("left",{})
        for nm in ("Hardening chamber (13 positions)","Controlled cooling chamber (1)","Tempering chamber (12 positions)"):
            cell(ws,f"B{r}",nm); v=px.get(nm,("","",""))
            cell(ws,f"C{r}",v[0] if v[0]!="" else None,font(10,False,BLUE if v[0]!="" else "000000"),INPUT,CEN)
            mbox(ws,f"D{r}:F{r}",v[1] or None,font(10,False,BLUE if v[1] else "000000"),INPUT); mbox(ws,f"G{r}:I{r}",v[2] or None,font(10,False,BLUE if v[2] else "000000"),INPUT); r+=1
        cell(ws,f"B{r}","Trays with product left inside, total",font(9,True),SEC); cell(ws,f"C{r}",f"=SUM(C{p0}:C{r-1})",font(10,True),CALC,CEN); LEFT=f"$C${r}"; r+=1
        cell(ws,f"B{r}","Product left inside: disposition recorded?",font(9,True),SEC)
        mbox(ws,f"C{r}:I{r}",f'=IF({LEFT}=0,"No product left inside",IF(COUNTA(G{p0}:G{p0+2})>=COUNTIF(C{p0}:C{p0+2},">0"),"Disposition recorded for the trays left inside","Record the disposition of the trays left inside (hold for evaluation; hardness before release)"))',font(10,True),CALC); LEFTS=f"$C${r}"; r+=2
    # summary
    mbox(ws,f"A{r}:I{r}","Summary (calculated)",font(10,True,"FFFFFF"),HEAD); r+=1
    def srow(lbl,frm,bold=False):
        nonlocal r
        mbox(ws,f"A{r}:D{r}",lbl,font(9,True),SEC); c=mbox(ws,f"E{r}:I{r}",frm,font(11 if bold else 10,True),CALC); r+=1; return f"$E${r-1}"
    APP=f'{rng("I")},"Yes"'
    tot=srow("Check items that apply to this furnace",f'=COUNTIF({rng("I")},"Yes")')
    done=srow("Answered (OK / Not OK / N/A)",f'=COUNTIFS({APP},{rng("E")},"<>")')
    nok=srow("Critical items Not OK",f'=COUNTIFS({APP},{rng("C")},"Yes",{rng("E")},"Not OK")')
    nok2=srow("Other items Not OK",f'=COUNTIFS({APP},{rng("C")},"No",{rng("E")},"Not OK")')
    opn=srow("Critical items not answered",f'=COUNTIFS({APP},{rng("C")},"Yes",{rng("E")},"")')
    if su:
        zo=srow("Zones out of tolerance when declared ready",f'=COUNTIF(H{z0}:H{z1},"OUT")')
        zm=srow("Zones without readings",f'=COUNTBLANK(H{z0}:H{z1})')
        res=srow("RESULT",f'=IF({nok}>0,"NOT RELEASED – critical item Not OK: start the reaction plan in CALOR (Alarms)",IF(OR({opn}>0,{zm}>0,{SOAK}=""),"INCOMPLETE – answer every critical item, the zone readings and the times",IF({zo}>0,"NOT RELEASED – zone out of tolerance",IF(LEFT({SOAKS},5)="SHORT","NOT RELEASED – soak too short",IF(LEFT({FTS},4)="LATE","RELEASED – first-off test late: record why","RELEASED FOR PRODUCTION")))))',True)
    else:
        res=srow("RESULT",f'=IF({nok}>0,"SHUTDOWN NOT SAFE / NOT COMPLETE – critical item Not OK",IF({opn}>0,"INCOMPLETE – answer every critical item",IF(AND({LEFT}>0,LEFT({LEFTS},11)="Record the "),"COMPLETE – but record the disposition of the product left inside","SHUTDOWN COMPLETE")))',True)
    rr=int(res[3:])
    ws.conditional_formatting.add(f"E{rr}",FormulaRule(formula=[f'OR(LEFT(E{rr},3)="NOT",LEFT(E{rr},8)="SHUTDOWN N")'],fill=red,font=Font(name=F,bold=True,color="9C0006")))
    ws.conditional_formatting.add(f"E{rr}",FormulaRule(formula=[f'OR(LEFT(E{rr},10)="RELEASED F",E{rr}="SHUTDOWN COMPLETE")'],fill=PatternFill("solid",fgColor="D5EFD9"),font=Font(name=F,bold=True,color="1B5E20")))
    ws.conditional_formatting.add(f"E{rr}",FormulaRule(formula=[f'OR(LEFT(E{rr},10)="INCOMPLETE",LEFT(E{rr},10)="RELEASED –",LEFT(E{rr},10)="COMPLETE –")'],fill=PatternFill("solid",fgColor="FCE8C3"),font=Font(name=F,bold=True,color="7A4A00")))
    r+=1
    # remarks + signatures
    mbox(ws,f"A{r}:B{r}","Remarks / Not OK items: action taken",font(9,True),SEC); mbox(ws,f"C{r}:I{r}",ex.get("remarks"),font(10,False,BLUE if ex.get("remarks") else "000000"),INPUT,WRAP); ws.row_dimensions[r].height=48; r+=2
    for i,h in enumerate(["","Sign-off","Name","","Signature","","Date / time","",""]):
        if h: cell(ws,f"{'ABCDEFGHI'[i]}{r}",h,font(9,True),SEC,CEN)
    ws.merge_cells(f"C{r}:D{r}"); ws.merge_cells(f"E{r}:F{r}"); ws.merge_cells(f"G{r}:I{r}"); r+=1
    sg=ex.get("sign",{})
    for who in ("Furnace operator","Shift supervisor","HT Manager (release)" if su else "HT Manager / Maintenance (handover)"):
        cell(ws,f"B{r}",who,font(9,True)); v=sg.get(who,"")
        mbox(ws,f"C{r}:D{r}",v or None,font(10,False,BLUE if v else "000000"),INPUT); mbox(ws,f"E{r}:F{r}",None,font(10),INPUT); mbox(ws,f"G{r}:I{r}",None,font(10),INPUT)
        ws.row_dimensions[r].height=26; r+=1
    # print setup
    ws.print_area=f"A1:I{r-1}"; ws.page_setup.orientation="portrait"; ws.page_setup.paperSize=ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0; ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.print_title_rows="1:3"; ws.page_margins.left=ws.page_margins.right=0.4; ws.page_margins.top=ws.page_margins.bottom=0.5
    ws.oddFooter.left.text="&8Form "+("F-HT-02A" if su else "F-HT-03A")+" Rev 0 · &A"; ws.oddFooter.right.text="&8Page &P of &N"
    ws.freeze_panes="A5"; ws.sheet_view.showGridLines=False
    return ws

wb=Workbook()
ins=wb.active; ins.title="Instructions"
setw(ins,{"A":3,"B":110})
L=[("Furnace start-up and shutdown checksheets",font(15,True)),
 ("Mass Wire and Steels Pvt. Ltd. · Bhiwadi · pusher furnaces F-01 … F-04",font(10,False,"555555")),
 ("",None),("Sheets",font(11,True)),
 ("Start-up – form F-HT-02A: one sheet per start-up (cold start, restart after power failure, maintenance, an alarm stop or a holiday).",font(10)),
 ("Shutdown – form F-HT-03A: one sheet per shutdown (planned, maintenance, holiday, emergency, breakdown).",font(10)),
 ("Example start-up / Example shutdown – the same forms filled in for F-02, to show the expected entries.",font(10)),
 ("",None),("How to fill in",font(11,True)),
 ("• Fill only the yellow cells. Grey cells are calculated; dark-blue headings and labels are fixed.",font(10)),
 ("• Choose the furnace first: items for another furnace (water quench F-01, oil quench F-02) then show grey and 'N/A for F-0x' in the Applicable column. F-03 and F-04 have no quench tank.",font(10)),
 ("• Result: choose OK, Not OK or N/A from the list. Write the value or reference where the item asks for one (gas pressure, quench temperature, specification no., test no.).",font(10)),
 ("• Critical = Yes: a Not OK result stops the release (start-up) or makes the shutdown unsafe; start a reaction plan on the CALOR Alarms page.",font(10)),
 ("• Times: type date and time as 21-10-2026 14:30. The soak (zones in tolerance → first tray) and the first-off test delay are calculated.",font(10)),
 ("• Zone readings: the set value and both wall thermocouples of every zone when the furnace is declared ready; each must be within the tolerance in Settings.",font(10)),
 ("• To use a sheet again, copy it (right-click the tab → Move or Copy → Create a copy) and clear the yellow cells, or print it blank and fill in by hand.",font(10)),
 ("",None),("Settings and where the items come from",font(11,True)),
 ("• Tolerance ±10 °C, minimum soak 30 min and first-off test within 8 h are the CALOR defaults (site rules, not CQI-9 figures). Change them in the blue Settings cells to your work instructions.",font(10)),
 ("• The check items are the start-up / shut-down checklists of CALOR (Process → Alarms) adapted to this plant: controlled cooling chamber on every furnace, F-01 water quench, F-02 oil quench, quench agitation checked by eye (not monitored), HK trays max. 200 kg.",font(10)),
 ("• Confirm the items against the furnace manufacturer's manual and the burner management / gas safety rules before use, and keep the form under document control (WI-HT-02 start-up, WI-HT-03 shut-down).",font(10)),
 ("• CQI-9 HTSA 1.15 asks for documented start-up and shut-down procedures; these forms are the records of following them. Enter each start-up / shutdown also in CALOR (Process → Alarms) so the Data review (1.9) and product traceability see it.",font(10)),
]
for i,(t,f) in enumerate(L,1):
    c=ins.cell(row=i,column=2,value=t or None); c.font=f or font(); c.alignment=Alignment(wrap_text=True,vertical="top")
ins.sheet_view.showGridLines=False
c=ins.cell(row=len(L)+2,column=2,value="Fill-in cell (yellow)"); c.fill=INPUT; c.font=font(); c.border=BOX
c=ins.cell(row=len(L)+3,column=2,value="Calculated cell (grey)"); c.fill=CALC; c.font=font(); c.border=BOX

build(wb.create_sheet("Start-up"),"su")
build(wb.create_sheet("Shutdown"),"sd")

d=datetime
exs={"furn":"F-02","date":d(2026,10,20),"shift":"A (06:00-14:00)","type":"Restart after maintenance","reason":"Burner 2R nozzle replaced (maintenance order MO-2231)",
 "res":{1:("OK","MO-2231 closed","R. Kumar",d(2026,10,20,6,5).time()),2:("OK","42 mbar","R. Kumar",d(2026,10,20,6,8).time()),3:("OK","","R. Kumar",d(2026,10,20,6,10).time()),4:("OK","","R. Kumar",d(2026,10,20,6,15).time()),
  5:("OK","","R. Kumar",d(2026,10,20,6,20).time()),6:("OK","SAT due 02-Dec-2026","S. Yadav",d(2026,10,20,6,22).time()),7:("OK","","S. Yadav",d(2026,10,20,6,25).time()),8:("OK","","R. Kumar",d(2026,10,20,6,25).time()),
  9:("OK","","R. Kumar",d(2026,10,20,6,40).time()),10:("OK","","R. Kumar",d(2026,10,20,6,45).time()),11:("N/A","","R. Kumar",d(2026,10,20,6,45).time()),12:("OK","62 °C","R. Kumar",d(2026,10,20,6,50).time()),
  13:("OK","","R. Kumar",d(2026,10,20,6,50).time()),14:("Not OK","Tray T2-11 cracked","S. Yadav",d(2026,10,20,7,0).time()),15:("OK","","R. Kumar",d(2026,10,20,9,30).time()),16:("OK","PS-0001 rev 1","S. Yadav",d(2026,10,20,9,30).time()),
  17:("OK","","S. Yadav",d(2026,10,20,10,15).time()),18:("OK","MT-0021","A. Meena",d(2026,10,20,13,5).time()),19:("OK","OP-0012","S. Yadav",d(2026,10,20,13,20).time())},
 "lit":d(2026,10,20,6,30),"ready":d(2026,10,20,9,40),"push":d(2026,10,20,10,15),"test":d(2026,10,20,13,0),
 "zones":{("Hardening",1):(720,724,718),("Hardening",2):(870,866,872),("Hardening",3):(870,873,868),("Tempering",1):(560,557,562),("Tempering",2):(580,584,578),("Tempering",3):(580,579,583)},
 "remarks":"Item 14: tray T2-11 cracked at a corner node – taken out, quarantined, inspection recorded in CALOR → Load build. Not critical: start-up released.",
 "sign":{"Furnace operator":"R. Kumar","Shift supervisor":"S. Yadav","HT Manager (release)":"V. Singh"}}
build(wb.create_sheet("Example start-up"),"su",exs)
exd={"furn":"F-02","date":d(2026,10,24),"shift":"C (22:00-06:00)","type":"Holiday / weekend","reason":"Diwali shutdown, 2 days (no orders)",
 "res":{1:("OK","Plan from Production Manager","R. Kumar",d(2026,10,24,22,0).time()),2:("OK","Tray 58 (MW-4512 F2-2610-09)","R. Kumar",d(2026,10,24,22,5).time()),3:("OK","25 empty trays","R. Kumar",d(2026,10,24,22,10).time()),
  4:("OK","Line empty","S. Yadav",d(2026,10,25,2,15).time()),5:("OK","","S. Yadav",d(2026,10,25,2,20).time()),6:("OK","Locked, tag 14","S. Yadav",d(2026,10,25,2,25).time()),7:("OK","","S. Yadav",d(2026,10,25,2,25).time()),
  8:("OK","","R. Kumar",d(2026,10,25,2,30).time()),9:("N/A","","R. Kumar",d(2026,10,25,2,30).time()),10:("OK","Cooler off, cover on","R. Kumar",d(2026,10,25,2,35).time()),11:("OK","","R. Kumar",d(2026,10,25,2,40).time()),
  12:("OK","","S. Yadav",d(2026,10,25,2,45).time()),13:("OK","No parts found","Maintenance fitter",d(2026,10,26,9,0).time()),14:("OK","OP-0015","S. Yadav",d(2026,10,25,2,50).time())},
 "stop":d(2026,10,24,22,5),"lastout":d(2026,10,25,2,10),"off":d(2026,10,25,2,20),"gas":d(2026,10,25,2,25),
 "left":{"Hardening chamber (13 positions)":(0,"",""),"Controlled cooling chamber (1)":(0,"",""),"Tempering chamber (12 positions)":(0,"","")},
 "remarks":"Line run out with empty trays before stopping; no product left inside.",
 "sign":{"Furnace operator":"R. Kumar","Shift supervisor":"S. Yadav","HT Manager / Maintenance (handover)":"Maintenance Head"}}
build(wb.create_sheet("Example shutdown"),"sd",exd)
for n in ("Example start-up","Example shutdown"): wb[n].sheet_properties.tabColor="9AA5A8"
for n in ("Start-up","Shutdown"): wb[n].sheet_properties.tabColor="1F4E5A"
out="/home/user/grain-boundary-project/documents/Furnace_Startup_Shutdown_Checksheet.xlsx"
wb.save(out); print(out)
